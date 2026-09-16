from __future__ import annotations

import errno
import hashlib
import json
import os
import re
import shlex
import shutil
import stat
import subprocess
import tempfile
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from .config import SandboxConfig


RESOURCE_READER_SCRIPT = r'''import datetime,json,sys
from pathlib import Path

kind, relative, raw_offset, raw_limit, raw_max = sys.argv[1:]
root = Path('/workspace').resolve()
target = (root / relative).resolve()
target.relative_to(root)
offset = max(0, int(raw_offset))
limit = int(raw_limit)
max_output = max(1, int(raw_max))

def safe(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime.date, datetime.datetime, datetime.time)):
        return value.isoformat()
    return str(value)

if kind == 'pdf':
    from pypdf import PdfReader
    reader = PdfReader(str(target))
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        pages.append(f'\n--- page {number} ---\n' + (page.extract_text() or ''))
    text = ''.join(pages)
    cap = max_output if limit < 0 else min(max_output, max(0, limit))
    selected = text[offset:offset + cap]
    result = {'resource': {
        'path': relative, 'type': 'application/pdf',
        'metadata': {'size': target.stat().st_size, 'pages': len(reader.pages), 'characters': len(text), 'encrypted': bool(reader.is_encrypted)},
        'representations': [{'kind': 'text', 'offset': offset, 'text': selected, 'truncated': offset + len(selected) < len(text)}],
    }}
elif kind == 'xlsx':
    from openpyxl import load_workbook
    workbook = load_workbook(target, read_only=True, data_only=True)
    row_limit = 3 if limit < 0 else min(100, max(0, limit))
    sheets = []
    previews = []
    for sheet in workbook.worksheets:
        sheets.append({'name': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column})
        rows = []
        if row_limit:
            for row in sheet.iter_rows(min_row=offset + 1, max_row=offset + row_limit, max_col=min(sheet.max_column or 1, 50), values_only=True):
                rows.append([safe(value) for value in row])
        previews.append({'kind': 'sheet_preview', 'sheet': sheet.title, 'row_offset': offset, 'rows': rows, 'truncated': offset + len(rows) < (sheet.max_row or 0)})
    result = {'resource': {
        'path': relative, 'type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'metadata': {'size': target.stat().st_size, 'sheets': sheets},
        'representations': previews,
    }}
else:
    raise ValueError('Unsupported resource kind')
print(json.dumps(result, ensure_ascii=False))
'''

HTTP_READER_SCRIPT = r'''import hashlib,json,sys,urllib.request
from urllib.parse import urlsplit

url, raw_offset, raw_limit, raw_max = sys.argv[1:]
parsed = urlsplit(url)
if parsed.scheme.lower() not in {'http', 'https'} or not parsed.hostname:
    raise ValueError('Only absolute http/https URLs are supported')
if parsed.username or parsed.password or (parsed.port not in {None, 80, 443}):
    raise ValueError('URL credentials and non-standard ports are forbidden')
offset = max(0, int(raw_offset))
limit = int(raw_limit)
max_output = max(1, int(raw_max))
request = urllib.request.Request(url, headers={
    'User-Agent': 'AIOS-ResourceAdapter/1.0 (+governed-read)',
    'Accept': 'text/html,application/xhtml+xml,text/plain,application/json;q=0.9,*/*;q=0.5',
    'Accept-Language': 'zh-CN,zh;q=0.9,en;q=0.5',
})
with urllib.request.urlopen(request, timeout=20) as response:
    raw = response.read(max_output + offset + 1)
    status = int(getattr(response, 'status', 200))
    final_url = response.geturl()
    content_type = response.headers.get_content_type()
    charset = response.headers.get_content_charset() or 'utf-8'
text = raw.decode(charset, errors='replace')
cap = max_output if limit < 0 else min(max_output, max(0, limit))
selected = text[offset:offset + cap]
result = {'resource': {
    'path': url,
    'type': content_type,
    'metadata': {
        'status': status,
        'requested_url': url,
        'final_url': final_url,
        'source_domain': parsed.hostname.lower().removeprefix('www.'),
        'bytes_observed': len(raw),
        'content_digest': hashlib.sha256(raw).hexdigest(),
        'characters_observed': len(text),
    },
    'representations': [{
        'kind': 'text', 'offset': offset, 'text': selected,
        'truncated': offset + len(selected) < len(text) or len(raw) > max_output + offset,
    }],
}}
print(json.dumps(result, ensure_ascii=False))
'''


class SandboxUnavailable(RuntimeError):
    pass


class SandboxPolicyError(RuntimeError):
    pass


@dataclass(slots=True)
class SandboxSession:
    task_id: int
    path: Path
    state_path: Path
    dependencies_path: Path


class DockerSandboxBroker:
    """Transactional workspace plus Docker-only command execution.

    There is deliberately no host-shell fallback. A missing Docker daemon is a
    capability failure, not permission to execute model-generated commands on
    the host.
    """

    def __init__(
        self, root: Path, config: SandboxConfig, skills_root: Path | None = None,
        *, network_enabled: bool = False, self_versions: Any = None,
    ):
        self.root = root.resolve()
        self.config = config
        self.skills_root = skills_root.resolve() if skills_root is not None else None
        self.network_enabled = bool(network_enabled)
        self.self_versions = self_versions
        self.session: SandboxSession | None = None
        self.dependencies_root = (self.root / "task_dependencies").resolve()
        # A recoverable task failure must not publish an unverified workspace,
        # but it also must not erase hours of valid intermediate work.  Keep a
        # task-private snapshot outside the ephemeral Docker session tree.  A
        # later attempt resumes from it and only ``commit`` publishes it.
        self.task_workspaces_root = (self.root / "task_workspaces").resolve()
        self.last_prepare_resumed = False
        self.scientific_environment = (self.root / "environments" / "scientific-py312-v1").resolve()
        self._health_state = "unknown"
        self._health_checked_at = 0.0
        self._health_probe_count = 0
        self._health_error: str | None = None

    def available(self, *, force: bool = False) -> bool:
        ttl = max(0, int(self.config.health_ttl_seconds))
        if (
            not force
            and self._health_state != "unknown"
            and time.monotonic() - self._health_checked_at <= ttl
        ):
            return self._health_state == "ready"
        return self._probe_health()

    def _probe_health(self) -> bool:
        self._health_probe_count += 1
        if self.config.backend != "docker" or shutil.which("docker") is None:
            self._health_state = "unavailable"
            self._health_checked_at = time.monotonic()
            self._health_error = "docker executable unavailable"
            return False
        try:
            result = subprocess.run(
                ["docker", "info", "--format", "{{.ServerVersion}}"],
                capture_output=True,
                text=True, encoding="utf-8", errors="replace",
                timeout=8,
                check=False,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            self._health_state = "unavailable"
            self._health_checked_at = time.monotonic()
            self._health_error = f"{type(exc).__name__}: {exc}"
            return False
        ready = result.returncode == 0
        self._health_state = "ready" if ready else "unavailable"
        self._health_checked_at = time.monotonic()
        self._health_error = None if ready else str(result.stderr or result.stdout)[-1000:]
        return ready

    def invalidate_health(self, reason: str = "explicit_invalidation") -> None:
        self._health_state = "unknown"
        self._health_checked_at = 0.0
        self._health_error = reason

    @property
    def health_probe_count(self) -> int:
        return self._health_probe_count

    def health_snapshot(self) -> dict[str, Any]:
        return {
            "state": self._health_state,
            "checked_at_monotonic": self._health_checked_at,
            "probe_count": self._health_probe_count,
            "error": self._health_error,
            "session_task_id": self.session.task_id if self.session is not None else None,
        }

    def describe_skill_invocation(self, command: str) -> dict[str, Any] | None:
        parsed = _parse_skill_command(command)
        if parsed is None or parsed["operation"] != "run" or self.skills_root is None:
            return None
        payload = parsed.get("input", {})
        canonical = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        input_metadata = {
            "input_digest": hashlib.sha256(canonical.encode("utf-8")).hexdigest(),
            "input_keys": sorted(str(key) for key in payload),
        }
        package = self.skills_root / "active" / parsed["name"]
        manifest_path = package / "manifest.json"
        if not manifest_path.is_file():
            return {
                **parsed,
                "version": "unknown",
                "required_capabilities": [],
                "manifest_found": False,
                **input_metadata,
            }
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        return {
            **parsed,
            "version": str(manifest.get("version", "unknown")),
            "required_capabilities": [str(item) for item in manifest.get("required_capabilities", [])],
            **input_metadata,
            "manifest_found": True,
        }

    def prepare(self, task_id: int, workspace: Path) -> Path:
        session_root = (self.root / f"task_{task_id}").resolve()
        if session_root.parent != self.root:
            raise ValueError("Invalid sandbox session path")
        if session_root.exists():
            self._remove_tree(session_root)
        session_root.mkdir(parents=True)
        snapshot = session_root / "workspace"
        state_path = session_root / "state"
        dependencies_path = (self.dependencies_root / f"task_{task_id}").resolve()
        task_workspace = self._task_workspace_path(task_id)
        source = task_workspace if task_workspace.is_dir() else workspace.resolve()
        self.last_prepare_resumed = source == task_workspace
        shutil.copytree(source, snapshot, dirs_exist_ok=True)
        state_path.mkdir()
        dependencies_path.mkdir(parents=True, exist_ok=True)
        self.session = SandboxSession(task_id, snapshot, state_path, dependencies_path)
        self.invalidate_health("session_created")
        self.available(force=True)
        return snapshot

    def run(self, command: str, timeout_seconds: int | None = None) -> dict[str, Any]:
        if self.session is None:
            raise SandboxUnavailable("No active sandbox session")
        self._validate_command_scope(command)
        if not self.available():
            raise SandboxUnavailable("Docker sandbox is unavailable; host execution is forbidden")
        before = self._manifest(self.session.path)
        self_before = (
            self._manifest(self.self_versions.self_path)
            if self.self_versions is not None and self.self_versions.exposed else {}
        )
        timeout = self._effective_timeout(timeout_seconds)
        args = [
            "docker", "run", "--rm", "--network", self.network_mode, "--read-only",
            "--cap-drop", "ALL", "--security-opt", "no-new-privileges",
            "--memory", f"{self.config.memory_mb}m", "--cpus", str(self.config.cpus),
            "--pids-limit", str(self.config.pids_limit),
            "--mount", f"type=bind,src={self.session.path},dst=/workspace",
            "--mount", f"type=bind,src={self.session.state_path},dst=/aios-state,readonly",
            "--mount", f"type=bind,src={self.session.dependencies_path},dst=/deps",
            "--env", "PYTHONPATH=/deps:/opt/aios-scientific", "--env", "PIP_TARGET=/deps",
            "--tmpfs", "/tmp:rw,nosuid,size=128m",
        ]
        if self.scientific_environment.exists():
            args.extend(["--mount", f"type=bind,src={self.scientific_environment},dst=/opt/aios-scientific,readonly"])
        if self.skills_root is not None and self.skills_root.exists():
            args.extend(["--mount", f"type=bind,src={self.skills_root},dst=/skills,readonly"])
        if self.self_versions is not None and self.self_versions.exposed:
            self_mount = f"type=bind,src={self.self_versions.self_path},dst=/self"
            if not self.self_versions.writable:
                self_mount += ",readonly"
            args.extend([
                "--mount", self_mount,
                "--mount", f"type=bind,src={self.self_versions.versions},dst=/self-history,readonly",
            ])
        args.extend(["--workdir", "/workspace", self.config.image, "bash", "-o", "pipefail", "-lc", command])
        try:
            result = subprocess.run(
                args,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=timeout,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise TimeoutError(f"Sandbox command exceeded {timeout}s") from exc
        after = self._manifest(self.session.path)
        changes = sorted(name for name in set(before) | set(after) if before.get(name) != after.get(name))
        if self.self_versions is not None and self.self_versions.exposed:
            self_after = self._manifest(self.self_versions.self_path)
            changes.extend(
                f"self:{name}" for name in sorted(set(self_before) | set(self_after))
                if self_before.get(name) != self_after.get(name)
            )
        return {
            "exit_code": result.returncode,
            "stdout": result.stdout,
            "stderr": result.stderr,
            "changes": changes,
        }

    def run_candidate(self, package: Path, command: str, timeout_seconds: int) -> dict[str, Any]:
        self._validate_command_scope(command)
        if not self.available():
            raise SandboxUnavailable("Docker sandbox is unavailable; candidate code cannot run on host")
        self.root.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="skill_benchmark_", dir=self.root) as temporary:
            root = Path(temporary)
            workspace = root / "workspace"
            state = root / "state"
            workspace.mkdir()
            state.mkdir()
            (state / "state.json").write_text(
                json.dumps({"tasks": [], "traces": [], "dead-letters": [], "memory": [], "capabilities": {}}, ensure_ascii=False),
                encoding="utf-8",
            )
            timeout = self._effective_timeout(timeout_seconds)
            args = [
                "docker", "run", "--rm", "--network", self.network_mode, "--read-only",
                "--cap-drop", "ALL", "--security-opt", "no-new-privileges",
                "--memory", f"{self.config.memory_mb}m", "--cpus", str(self.config.cpus),
                "--pids-limit", str(self.config.pids_limit),
                "--mount", f"type=bind,src={workspace},dst=/workspace",
                "--mount", f"type=bind,src={state},dst=/aios-state,readonly",
                "--mount", f"type=bind,src={package.resolve()},dst=/candidate,readonly",
                "--workdir", "/workspace", self.config.image, "bash", "-o", "pipefail", "-lc", command,
            ]
            try:
                result = subprocess.run(
                    args, capture_output=True, text=True, encoding="utf-8", errors="replace",
                    timeout=timeout, check=False,
                )
            except subprocess.TimeoutExpired as exc:
                raise TimeoutError(f"Skill benchmark exceeded {timeout}s") from exc
            return {"exit_code": result.returncode, "stdout": result.stdout, "stderr": result.stderr}

    def run_self_harness(
        self, version: str, stage: str, payload: dict[str, Any],
        *, timeout_seconds: int = 20,
    ) -> dict[str, Any]:
        """Execute the versioned Agent entrypoint without importing it into Host."""
        if self.self_versions is None:
            raise SandboxUnavailable("Mutable Self is not configured")
        if re.fullmatch(r"v\d{6}", version) is None:
            raise SandboxPolicyError("Invalid Self version")
        source = (self.self_versions.versions / version).resolve()
        if source.parent != self.self_versions.versions.resolve() or not source.is_dir():
            raise SandboxPolicyError(f"Unknown Self version: {version}")
        if stage not in {"before_model", "after_plan", "after_round"}:
            raise SandboxPolicyError(f"Unknown Self Agent lifecycle event: {stage}")
        agent_entrypoint = source / "agent" / "main.py"
        legacy_entrypoint = source / "harness" / "runner.py"
        if agent_entrypoint.is_file():
            entrypoint = "/self/agent/main.py"
            runtime_kind = "agent"
        elif legacy_entrypoint.is_file():
            entrypoint = "/self/harness/runner.py"
            runtime_kind = "legacy_harness"
        else:
            raise RuntimeError("Self version has no Agent or legacy Harness entrypoint")
        if not self.available():
            raise SandboxUnavailable("Docker sandbox is unavailable; Self Agent cannot run on Host")
        self.root.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="self_harness_", dir=self.root) as temporary:
            exchange = Path(temporary).resolve()
            input_path = exchange / "input.json"
            output_path = exchange / "output.json"
            input_path.write_text(
                json.dumps(payload, ensure_ascii=False, default=str), encoding="utf-8",
            )
            timeout = self._effective_timeout(timeout_seconds)
            args = [
                "docker", "run", "--rm", "--network", "none", "--read-only",
                "--cap-drop", "ALL", "--security-opt", "no-new-privileges",
                "--memory", f"{self.config.memory_mb}m", "--cpus", str(self.config.cpus),
                "--pids-limit", str(self.config.pids_limit),
                "--mount", f"type=bind,src={source},dst=/self,readonly",
                "--mount", f"type=bind,src={exchange},dst=/exchange",
                "--tmpfs", "/tmp:rw,nosuid,size=32m",
                self.config.image, "python", entrypoint, stage,
                "/exchange/input.json", "/exchange/output.json",
            ]
            try:
                result = subprocess.run(
                    args, capture_output=True, text=True, encoding="utf-8", errors="replace",
                    timeout=timeout, check=False,
                )
            except subprocess.TimeoutExpired as exc:
                raise TimeoutError(f"Self Agent lifecycle event exceeded {timeout}s") from exc
            if result.returncode != 0:
                detail = str(result.stderr or result.stdout)[-4000:]
                raise RuntimeError(f"Self Agent {stage} failed: {detail}")
            if not output_path.is_file():
                raise RuntimeError(f"Self Agent {stage} produced no output")
            if output_path.stat().st_size > 2_000_000:
                raise RuntimeError(f"Self Agent {stage} output exceeds 2000000 bytes")
            try:
                output = json.loads(output_path.read_text(encoding="utf-8"))
            except json.JSONDecodeError as exc:
                raise RuntimeError(f"Self Agent {stage} returned invalid JSON") from exc
            if not isinstance(output, dict):
                raise RuntimeError(f"Self Agent {stage} must return an object")
            return {
                "output": output, "version": version, "stage": stage,
                "runtime_kind": runtime_kind,
                "entrypoint": entrypoint.removeprefix("/self/"),
                "stdout": result.stdout[-4000:], "stderr": result.stderr[-4000:],
            }

    def read_resource(
        self, relative_path: str, *, kind: str, offset: int = 0,
        limit: int | None = None, max_output_bytes: int = 1_048_576,
    ) -> dict[str, Any]:
        """Run a fixed resource adapter; no model-generated shell is involved."""
        if self.session is None:
            raise SandboxUnavailable("No active sandbox session")
        packages = {"pdf": ("pypdf", "pypdf==5.4.0"), "xlsx": ("openpyxl", "openpyxl==3.1.5")}
        if kind not in packages:
            raise SandboxPolicyError(f"Unsupported resource adapter: {kind}")
        path = Path(relative_path)
        if path.is_absolute() or ".." in path.parts:
            raise SandboxPolicyError("Resource adapter path must be workspace-relative")
        if not self.available():
            self.invalidate_health("adapter_preflight_unavailable")
            if not self.available(force=True):
                raise SandboxUnavailable("Docker sandbox is unavailable; resource adapters cannot run on host")

        module, requirement = packages[kind]
        if not (self.session.dependencies_path / module).exists() and not (self.scientific_environment / module).exists():
            install = self.run(
                f"python -m pip install --disable-pip-version-check --no-input --target /deps {requirement}"
            )
            if install.get("exit_code") != 0:
                raise SandboxUnavailable(
                    "Resource adapter dependency installation failed: "
                    + str(install.get("stderr") or install.get("stdout") or "unknown error")[-2000:]
                )

        script = self.session.state_path / "resource_reader.py"
        script.write_text(RESOURCE_READER_SCRIPT, encoding="utf-8")
        args = [
            "docker", "run", "--rm", "--network", "none", "--read-only",
            "--cap-drop", "ALL", "--security-opt", "no-new-privileges",
            "--memory", f"{self.config.memory_mb}m", "--cpus", str(self.config.cpus),
            "--pids-limit", str(self.config.pids_limit),
            "--mount", f"type=bind,src={self.session.path},dst=/workspace,readonly",
            "--mount", f"type=bind,src={self.session.state_path},dst=/aios-state,readonly",
            "--mount", f"type=bind,src={self.session.dependencies_path},dst=/deps,readonly",
            "--env", "PYTHONPATH=/deps:/opt/aios-scientific", "--tmpfs", "/tmp:rw,nosuid,size=128m",
            "--workdir", "/workspace", self.config.image,
            "python", "/aios-state/resource_reader.py", kind, path.as_posix(),
            str(max(0, offset)), str(-1 if limit is None else limit), str(max_output_bytes),
        ]
        if self.scientific_environment.exists():
            insert_at = args.index("--workdir")
            args[insert_at:insert_at] = ["--mount", f"type=bind,src={self.scientific_environment},dst=/opt/aios-scientific,readonly"]
        for attempt in (1, 2):
            try:
                result = subprocess.run(
                    args, capture_output=True, text=True, encoding="utf-8", errors="replace",
                    timeout=self.config.default_timeout_seconds, check=False,
                )
            except subprocess.TimeoutExpired as exc:
                raise TimeoutError(f"Resource adapter exceeded {self.config.default_timeout_seconds}s") from exc
            output = {
                "exit_code": result.returncode, "stdout": result.stdout, "stderr": result.stderr,
                "attempts": attempt, "retry_count": attempt - 1,
                "transient_recovered": attempt > 1 and result.returncode == 0,
            }
            if result.returncode == 0 or attempt == 2 or not self._transient_docker_failure(output):
                return output
            self.invalidate_health("transient_adapter_failure")
            if not self.available(force=True):
                return output
        raise AssertionError("unreachable adapter retry state")

    def read_http(
        self, url: str, *, offset: int = 0, limit: int | None = None,
        max_output_bytes: int = 12_000,
    ) -> dict[str, Any]:
        """Read an HTTP resource with fixed host-owned code behind `read(URL)`."""
        if self.session is None:
            raise SandboxUnavailable("No active sandbox session")
        if not self.network_enabled:
            raise SandboxPolicyError("External network authority is not granted")
        from urllib.parse import urlsplit
        parsed = urlsplit(str(url).strip())
        if (
            parsed.scheme.casefold() not in {"http", "https"}
            or not parsed.hostname
            or parsed.username is not None
            or parsed.password is not None
            or parsed.port not in {None, 80, 443}
        ):
            raise SandboxPolicyError("HTTP adapter accepts absolute http/https URLs without credentials or non-standard ports")
        if not self.available():
            raise SandboxUnavailable("Docker sandbox is unavailable; HTTP resources cannot be read")
        script = self.session.state_path / "http_reader.py"
        script.write_text(HTTP_READER_SCRIPT, encoding="utf-8")
        args = [
            "docker", "run", "--rm", "--network", self.network_mode, "--read-only",
            "--cap-drop", "ALL", "--security-opt", "no-new-privileges",
            "--memory", f"{self.config.memory_mb}m", "--cpus", str(self.config.cpus),
            "--pids-limit", str(self.config.pids_limit),
            "--mount", f"type=bind,src={self.session.state_path},dst=/aios-state,readonly",
            "--tmpfs", "/tmp:rw,nosuid,size=32m", self.config.image,
            "python", "/aios-state/http_reader.py", str(url), str(max(0, offset)),
            str(-1 if limit is None else limit), str(max(1, max_output_bytes)),
        ]
        for attempt in (1, 2):
            try:
                result = subprocess.run(
                    args, capture_output=True, text=True, encoding="utf-8", errors="replace",
                    timeout=self.config.default_timeout_seconds, check=False,
                )
            except subprocess.TimeoutExpired as exc:
                raise TimeoutError(f"HTTP resource adapter exceeded {self.config.default_timeout_seconds}s") from exc
            output = {
                "exit_code": result.returncode,
                "stdout": result.stdout,
                "stderr": result.stderr,
                "attempts": attempt,
                "retry_count": attempt - 1,
                "transient_recovered": attempt > 1 and result.returncode == 0,
            }
            if result.returncode == 0 or attempt == 2 or not self._transient_http_failure(output):
                return output
        raise AssertionError("unreachable HTTP adapter retry state")

    @staticmethod
    def _transient_http_failure(result: dict[str, Any]) -> bool:
        text = f"{result.get('stderr', '')} {result.get('stdout', '')}".casefold()
        return any(marker in text for marker in (
            "temporary failure in name resolution",
            "name or service not known",
            "connection reset",
            "connection refused",
            "timed out",
            "remote end closed connection",
        ))

    @staticmethod
    def _transient_docker_failure(result: dict[str, Any]) -> bool:
        if int(result.get("exit_code", 0)) not in {125, 126, 127}:
            return False
        text = f"{result.get('stderr', '')} {result.get('stdout', '')}".casefold()
        return any(marker in text for marker in (
            "cannot connect to the docker daemon", "error during connect",
            "connection refused", "docker daemon", "context deadline exceeded",
            "docker desktop", "the system cannot find the file specified",
        ))

    def ensure_scientific_environment(self) -> dict[str, Any]:
        """Provision the governed scientific stack once, then mount it read-only for tasks."""
        started = time.perf_counter()
        sentinel = self.scientific_environment / ".aios-environment.json"
        if sentinel.is_file():
            return {"ready": True, "cache_hit": True, "latency_ms": (time.perf_counter() - started) * 1000}
        if not self.network_enabled:
            return {"ready": False, "cache_hit": False, "error": "Network authority is required to provision the scientific environment"}
        if not self.available():
            return {"ready": False, "cache_hit": False, "error": "Docker sandbox is unavailable"}
        self.scientific_environment.mkdir(parents=True, exist_ok=True)
        requirements = [
            "numpy==2.2.6", "pandas==2.2.3", "scipy==1.15.3",
            "statsmodels==0.14.4", "openpyxl==3.1.5", "pypdf==5.4.0",
        ]
        args = [
            "docker", "run", "--rm", "--network", self.network_mode,
            "--cap-drop", "ALL", "--security-opt", "no-new-privileges",
            "--memory", f"{self.config.memory_mb}m", "--cpus", str(self.config.cpus),
            "--pids-limit", str(self.config.pids_limit),
            "--mount", f"type=bind,src={self.scientific_environment},dst=/environment",
            "--tmpfs", "/tmp:rw,nosuid,size=512m", self.config.image,
            "python", "-m", "pip", "install", "--disable-pip-version-check", "--no-input",
            "--target", "/environment", *requirements,
        ]
        try:
            result = subprocess.run(
                args, capture_output=True, text=True, encoding="utf-8", errors="replace",
                timeout=self.config.max_timeout_seconds, check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise TimeoutError(f"Scientific environment provisioning exceeded {self.config.max_timeout_seconds}s") from exc
        latency = (time.perf_counter() - started) * 1000
        if result.returncode != 0:
            return {"ready": False, "cache_hit": False, "latency_ms": latency, "error": (result.stderr or result.stdout)[-4000:]}
        sentinel.write_text(json.dumps({"requirements": requirements}, ensure_ascii=False), encoding="utf-8")
        return {"ready": True, "cache_hit": False, "latency_ms": latency, "requirements": requirements}

    def purge_task_dependencies(self, task_id: int) -> None:
        path = (self.dependencies_root / f"task_{task_id}").resolve()
        if path.parent == self.dependencies_root and path.exists():
            self._remove_tree(path)

    def _effective_timeout(self, requested: int | None) -> int:
        value = self.config.default_timeout_seconds if requested is None else int(requested)
        return max(1, min(value, self.config.max_timeout_seconds))

    @property
    def network_mode(self) -> str:
        """Docker egress mode. `bridge` is intentionally unrestricted."""
        return "bridge" if self.network_enabled else "none"

    @staticmethod
    def _validate_command_scope(command: str) -> None:
        _parse_skill_command(command)
        broad_root_patterns = (
            r"(?:^|[;&|]\s*)cd\s+/(?:\s|[;&|]|$)",
            r"\bfind\s+/(?:\s|$)",
            r"\bls\s+(?:-[A-Za-z]+\s+)?/(?:\s|[;&|]|$)",
            r"\b(?:grep|rg|du)\b[^;&|]*\s/(?:\s|[;&|]|$)",
        )
        if any(re.search(pattern, command) for pattern in broad_root_patterns):
            raise SandboxPolicyError(
                "Broad container-root access is forbidden; inspect only documented mounts"
            )

    def expose_read_only_state(self, state: dict[str, Any]) -> None:
        if self.session is None:
            raise SandboxUnavailable("No active sandbox session")
        internal = self.session.state_path
        (internal / "state.json").write_text(json.dumps(state, ensure_ascii=False, default=str), encoding="utf-8")
        (internal / "aiosctl.py").write_text(
            "import json,sys\n"
            "data=json.load(open('/aios-state/state.json',encoding='utf-8'))\n"
            "name=sys.argv[1] if len(sys.argv)>1 else 'capabilities'\n"
            "print(json.dumps(data.get(name,[]),ensure_ascii=False,indent=2))\n",
            encoding="utf-8",
        )

    def commit(self, workspace: Path) -> list[str]:
        if self.session is None:
            return []
        task_id = self.session.task_id
        destination = workspace.resolve()
        changed: list[str] = []
        current = self._manifest(destination)
        snapshot = self._manifest(self.session.path)
        for relative, digest in snapshot.items():
            if relative == ".aios" or relative.startswith(".aios/"):
                continue
            if current.get(relative) == digest:
                continue
            source = self.session.path / relative
            target = destination / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, target)
            changed.append(relative)
        self.discard()
        self.clear_task_workspace(task_id)
        return sorted(changed)

    def checkpoint_task_workspace(self, task_id: int | None = None) -> dict[str, Any]:
        """Persist the current task world without publishing it.

        This is a recovery primitive, not a successful commit.  It deliberately
        lives outside ``task_<id>`` because ``discard`` removes that ephemeral
        session after every failed cycle.
        """
        if self.session is None:
            return {"preserved": False, "reason": "no_active_session"}
        selected = self.session.task_id if task_id is None else int(task_id)
        if selected != self.session.task_id:
            raise SandboxPolicyError("Active sandbox belongs to another task")
        destination = self._task_workspace_path(selected)
        destination.parent.mkdir(parents=True, exist_ok=True)
        temporary = destination.parent / f".{destination.name}.{os.getpid()}.{time.time_ns()}.tmp"
        if temporary.exists():
            self._remove_tree(temporary)
        try:
            shutil.copytree(self.session.path, temporary)
            if destination.exists():
                self._remove_tree(destination)
            os.replace(temporary, destination)
        finally:
            if temporary.exists():
                self._remove_tree(temporary)
        manifest = self._manifest(destination)
        digest_source = json.dumps(manifest, sort_keys=True, separators=(",", ":"))
        self.discard()
        return {
            "preserved": True,
            "task_id": selected,
            "file_count": len(manifest),
            "manifest_digest": hashlib.sha256(digest_source.encode("utf-8")).hexdigest(),
        }

    def clear_task_workspace(self, task_id: int) -> None:
        path = self._task_workspace_path(task_id)
        if path.exists():
            self._remove_tree(path)

    def _task_workspace_path(self, task_id: int) -> Path:
        path = (self.task_workspaces_root / f"task_{int(task_id)}").resolve()
        if path.parent != self.task_workspaces_root:
            raise SandboxPolicyError("Invalid task workspace checkpoint path")
        return path

    def discard(self) -> None:
        if self.session is None:
            return
        session_root = self.session.path.parent
        self.session = None
        if session_root.parent == self.root and session_root.exists():
            self._remove_tree(session_root)

    @staticmethod
    def _remove_tree(path: Path) -> None:
        """Remove a managed tree despite Windows readonly and short-lived races."""
        def make_writable_and_retry(function: Any, target: str, exc_info: Any) -> None:
            error = exc_info[1]
            if not isinstance(error, PermissionError):
                raise error
            os.chmod(target, stat.S_IWRITE | stat.S_IREAD)
            function(target)

        for attempt in range(5):
            try:
                shutil.rmtree(path, onerror=make_writable_and_retry)
                return
            except FileNotFoundError:
                return
            except OSError as error:
                retryable = (
                    getattr(error, "winerror", None) in {5, 32, 145}
                    or error.errno in {errno.EACCES, errno.EBUSY, errno.ENOTEMPTY}
                )
                if not retryable or attempt == 4:
                    raise
                # Antivirus/indexers and recently stopped containers can briefly
                # retain directory entries. Clear attributes, yield, then retry the
                # whole managed tree instead of treating cleanup as a task failure.
                for target in [path, *path.rglob("*")]:
                    try:
                        os.chmod(target, stat.S_IWRITE | stat.S_IREAD)
                    except OSError:
                        pass
                time.sleep(0.05 * (2 ** attempt))

    @staticmethod
    def _manifest(root: Path) -> dict[str, str]:
        if not root.exists():
            return {}
        manifest: dict[str, str] = {}
        for path in root.rglob("*"):
            if path.is_file():
                relative = path.relative_to(root)
                if ".git" in relative.parts:
                    continue
                manifest[relative.as_posix()] = hashlib.sha256(path.read_bytes()).hexdigest()
        return manifest


def _is_json_object(value: str) -> bool:
    try:
        return isinstance(json.loads(value), dict)
    except (TypeError, json.JSONDecodeError):
        return False


def _parse_skill_command(command: str) -> dict[str, Any] | None:
    if "/skills" not in command:
        return None
    try:
        lexer = shlex.shlex(command, posix=True, punctuation_chars=";&|<>")
        lexer.whitespace_split = True
        lexer.commenters = ""
        tokens = list(lexer)
    except ValueError as exc:
        raise SandboxPolicyError("Invalid shell quoting in Skill invocation") from exc
    operators = {token for token in tokens if token and set(token) <= set(";&|<>")}
    skill_paths = [token for token in tokens if "/skills" in token]
    if operators or not skill_paths or any(path != "/skills/skill.py" for path in skill_paths):
        raise SandboxPolicyError(
            "Skill commands cannot be combined with shell operators or direct /skills access"
        )
    if len(tokens) < 3 or tokens[0] not in {"python", "python3"} or tokens[1] != "/skills/skill.py":
        raise SandboxPolicyError("Skills must be invoked through python /skills/skill.py")
    operation = tokens[2]
    name = tokens[3] if len(tokens) >= 4 else None
    payload: dict[str, Any] = {}
    valid = operation == "list" and len(tokens) == 3
    if operation == "show" and len(tokens) == 4 and isinstance(name, str):
        valid = re.fullmatch(r"[a-z][a-z0-9_]{1,63}", name) is not None
    if operation == "run" and isinstance(name, str) and re.fullmatch(r"[a-z][a-z0-9_]{1,63}", name):
        if len(tokens) == 4:
            valid = True
        elif len(tokens) == 6 and tokens[4] == "--input-json" and _is_json_object(tokens[5]):
            payload = json.loads(tokens[5])
            valid = True
    if not valid:
        raise SandboxPolicyError("Invalid or unsafe Skill dispatcher invocation")
    return {"operation": operation, "name": name, "input": payload}
